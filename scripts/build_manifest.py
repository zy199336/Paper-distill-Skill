from __future__ import annotations

import csv
import argparse
import re
import zipfile
from html import escape
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_PAPER_DIR = ROOT / "input_papers" / "pdfs"
DATA_DIR = ROOT / "data"
CSV_PATH = DATA_DIR / "paper_manifest.csv"
XLSX_PATH = DATA_DIR / "paper_manifest.xlsx"
RECOMMENDED_PDF_COUNT = 20

FIELDS = [
    "paper_id",
    "file_name",
    "file_path",
    "file_size_mb",
    "title",
    "authors",
    "year",
    "journal_or_venue",
    "doi",
    "domain",
    "subdomain",
    "paper_type",
    "method_type",
    "data_type",
    "scenario_type",
    "main_problem",
    "main_contribution_type",
    "has_clear_outline",
    "has_clear_problem_formulation",
    "has_good_writing_patterns",
    "has_good_experiment_design",
    "has_good_figures",
    "include_for_outline",
    "include_for_writing",
    "manual_priority",
    "extraction_status",
    "outline_extraction_status",
    "writing_extraction_status",
    "notes",
]

MANUAL_FIELDS = {
    "title",
    "authors",
    "year",
    "journal_or_venue",
    "doi",
    "domain",
    "subdomain",
    "paper_type",
    "method_type",
    "data_type",
    "scenario_type",
    "main_problem",
    "main_contribution_type",
    "has_clear_outline",
    "has_clear_problem_formulation",
    "has_good_writing_patterns",
    "has_good_experiment_design",
    "has_good_figures",
    "include_for_outline",
    "include_for_writing",
    "manual_priority",
    "extraction_status",
    "outline_extraction_status",
    "writing_extraction_status",
    "notes",
}


def read_existing_manifest() -> List[Dict[str, str]]:
    if not CSV_PATH.exists():
        return []
    with CSV_PATH.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        rows = []
        for row in reader:
            normalized = {field: (row.get(field) or "").strip() for field in FIELDS}
            rows.append(normalized)
        return rows


def write_csv(rows: List[Dict[str, str]]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    with CSV_PATH.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in FIELDS})


def write_xlsx(rows: List[Dict[str, str]]) -> Tuple[bool, str]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as exc:  # pragma: no cover - depends on local environment
        write_basic_xlsx(rows)
        return True, f"xlsx written with standard-library fallback because openpyxl is unavailable: {exc}"

    wb = Workbook()
    ws = wb.active
    ws.title = "paper_manifest"
    ws.append(FIELDS)
    for row in rows:
        ws.append([row.get(field, "") for field in FIELDS])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    widths = {
        "A": 12,
        "B": 34,
        "C": 48,
        "E": 42,
        "F": 32,
        "H": 24,
        "I": 26,
        "AC": 36,
    }
    for idx, field in enumerate(FIELDS, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = widths.get(letter, min(max(len(field) + 2, 14), 24))
    ws.freeze_panes = "A2"
    wb.save(XLSX_PATH)
    return True, "xlsx written"


def excel_cell(row: int, col: int) -> str:
    letters = ""
    while col:
        col, remainder = divmod(col - 1, 26)
        letters = chr(65 + remainder) + letters
    return f"{letters}{row}"


def write_basic_xlsx(rows: List[Dict[str, str]]) -> None:
    sheet_rows = [FIELDS] + [[row.get(field, "") for field in FIELDS] for row in rows]
    row_xml = []
    for row_idx, values in enumerate(sheet_rows, start=1):
        cells = []
        for col_idx, value in enumerate(values, start=1):
            ref = excel_cell(row_idx, col_idx)
            text = escape(str(value or ""))
            cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>')
        row_xml.append(f'<row r="{row_idx}">{"".join(cells)}</row>')

    worksheet = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>
  <sheetData>{"".join(row_xml)}</sheetData>
</worksheet>'''

    files = {
        "[Content_Types].xml": '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
  <Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
</Types>''',
        "_rels/.rels": '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
  <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>
</Relationships>''',
        "xl/workbook.xml": '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets><sheet name="paper_manifest" sheetId="1" r:id="rId1"/></sheets>
</workbook>''',
        "xl/_rels/workbook.xml.rels": '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
</Relationships>''',
        "xl/worksheets/sheet1.xml": worksheet,
        "docProps/core.xml": '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>paper_manifest</dc:title></cp:coreProperties>''',
        "docProps/app.xml": '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"><Application>Paper Skill Builder</Application></Properties>''',
    }
    with zipfile.ZipFile(XLSX_PATH, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for name, content in files.items():
            zf.writestr(name, content)


def decode_pdf_literal(value: str) -> str:
    value = value.strip()
    if value.startswith("<") and value.endswith(">") and not value.startswith("<<"):
        hex_text = re.sub(r"[^0-9A-Fa-f]", "", value[1:-1])
        try:
            raw = bytes.fromhex(hex_text)
            for encoding in ("utf-16-be", "utf-16", "utf-8", "latin-1"):
                try:
                    return raw.decode(encoding).replace("\x00", "").strip()
                except UnicodeDecodeError:
                    continue
        except ValueError:
            return ""
    if value.startswith("(") and value.endswith(")"):
        inner = value[1:-1]
        inner = inner.replace(r"\(", "(").replace(r"\)", ")").replace(r"\\", "\\")
        return inner.strip()
    return value.strip()


def extract_pdf_metadata(path: Path) -> Dict[str, str]:
    metadata = {
        "title": "",
        "authors": "",
        "year": "",
        "journal_or_venue": "unknown",
        "doi": "",
        "domain": "unknown",
        "subdomain": "unknown",
        "paper_type": "unknown",
        "method_type": "unknown",
        "data_type": "unknown",
        "scenario_type": "unknown",
    }
    try:
        blob = path.read_bytes()[:2_000_000]
    except OSError:
        return metadata

    text = blob.decode("latin-1", errors="ignore")

    for pdf_key, field in [("Title", "title"), ("Author", "authors")]:
        match = re.search(rf"/{pdf_key}\s*(\([^)]{{0,500}}\)|<[0-9A-Fa-f\s]{{2,1000}}>)", text, re.S)
        if match:
            value = decode_pdf_literal(match.group(1))
            if value and len(value) > 2:
                metadata[field] = value

    doi_match = re.search(r"\b10\.\d{4,9}/[-._;()/:A-Z0-9]+\b", text, re.I)
    if doi_match:
        metadata["doi"] = doi_match.group(0).rstrip(").,;")

    year_candidates = re.findall(r"\b(19[8-9]\d|20[0-3]\d)\b", text[:500_000])
    if year_candidates:
        metadata["year"] = sorted(year_candidates, reverse=True)[0]
    else:
        name_year = re.search(r"\b(19[8-9]\d|20[0-3]\d)\b", path.stem)
        if name_year:
            metadata["year"] = name_year.group(1)

    if not metadata["title"]:
        cleaned = re.sub(r"[_\-]+", " ", path.stem).strip()
        metadata["title"] = cleaned if cleaned else ""

    return metadata


def next_paper_id(used_ids: Iterable[str]) -> str:
    used_numbers = set()
    for paper_id in used_ids:
        match = re.fullmatch(r"PAPER_(\d{3})", paper_id or "")
        if match:
            used_numbers.add(int(match.group(1)))
    number = 1
    while number in used_numbers:
        number += 1
    return f"PAPER_{number:03d}"


def row_key(row: Dict[str, str]) -> str:
    return (row.get("file_path") or row.get("file_name") or "").replace("\\", "/").lower()


def resolve_paper_dir(value: str | None) -> Path:
    if not value:
        return DEFAULT_PAPER_DIR
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = (ROOT / path).resolve()
    return path


def find_pdfs(paper_dir: Path) -> List[Path]:
    if not paper_dir.exists():
        return []
    direct = sorted(paper_dir.glob("*.pdf"), key=lambda p: p.name.lower())
    if direct:
        return direct
    nested = sorted(paper_dir.rglob("*.pdf"), key=lambda p: p.as_posix().lower())
    return nested


def build_manifest(paper_dir: Path) -> int:
    paper_dir.mkdir(parents=True, exist_ok=True)
    pdfs = find_pdfs(paper_dir)
    if len(pdfs) > RECOMMENDED_PDF_COUNT:
        print(
            f"WARNING: found {len(pdfs)} PDFs in {paper_dir}. "
            f"For higher-quality skill generation, select a focused set of the most relevant, highest-quality papers "
            f"(recommended around {RECOMMENDED_PDF_COUNT}). Continuing anyway."
        )

    existing_rows = read_existing_manifest()
    existing_by_key = {row_key(row): row for row in existing_rows if row_key(row)}
    rows_by_id = {row.get("paper_id", ""): row for row in existing_rows if row.get("paper_id")}
    used_ids = set(rows_by_id)
    output_rows = list(existing_rows)
    added = 0
    updated = 0

    for pdf in pdfs:
        try:
            rel_path = pdf.relative_to(ROOT).as_posix()
        except ValueError:
            rel_path = pdf.resolve().as_posix()
        key = rel_path.lower()
        row = existing_by_key.get(key)
        if row is None:
            row = {field: "" for field in FIELDS}
            row["paper_id"] = next_paper_id(used_ids)
            used_ids.add(row["paper_id"])
            output_rows.append(row)
            added += 1

        metadata = extract_pdf_metadata(pdf)
        row["file_name"] = pdf.name
        row["file_path"] = rel_path
        row["file_size_mb"] = f"{pdf.stat().st_size / (1024 * 1024):.2f}"
        for field, value in metadata.items():
            if field in MANUAL_FIELDS and not row.get(field) and value:
                row[field] = value

        row["include_for_outline"] = row.get("include_for_outline") or "yes"
        row["include_for_writing"] = row.get("include_for_writing") or "yes"
        row["extraction_status"] = row.get("extraction_status") or "pending"
        row["outline_extraction_status"] = row.get("outline_extraction_status") or "pending"
        row["writing_extraction_status"] = row.get("writing_extraction_status") or "pending"
        updated += 1

    write_csv(output_rows)
    xlsx_ok, xlsx_message = write_xlsx(output_rows)

    print("Manifest build complete.")
    print(f"- Paper folder: {paper_dir}")
    print(f"- PDFs found: {len(pdfs)}")
    if len(pdfs) > RECOMMENDED_PDF_COUNT:
        print(f"- Recommendation: consider reducing to a focused set near {RECOMMENDED_PDF_COUNT} high-quality, highly relevant papers.")
    print(f"- Rows added: {added}")
    print(f"- PDF rows updated: {updated}")
    print(f"- CSV: {CSV_PATH}")
    if xlsx_ok:
        print(f"- XLSX: {XLSX_PATH}")
    else:
        print(f"- XLSX skipped: {xlsx_message}")
    return 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build or update paper_manifest from a folder of academic PDFs.")
    parser.add_argument(
        "--paper-dir",
        default=None,
        help="Folder containing PDF files. Defaults to input_papers/pdfs. If PDFs are not directly inside it, the script scans recursively.",
    )
    args = parser.parse_args()
    raise SystemExit(build_manifest(resolve_paper_dir(args.paper_dir)))
